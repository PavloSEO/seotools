"""Offline acceptance for the ``scan_export.v1`` scan-data export contract."""

from __future__ import annotations

import csv
import io
import json
import xml.etree.ElementTree as ET

import pytest

from seohead.storage import import_run
from seohead.storage import scan_export as export_module
from seohead.storage.scan_export import EXPORT_FORMAT_VERSION, XML_NAMESPACE, export_scan_data
from tests.test_scan_artifact import BUILD
from tests.test_scan_artifact import legacy_run as legacy_run


@pytest.fixture
def artifact(legacy_run, tmp_path):
    path = tmp_path / "scan.sqlite"
    import_run(legacy_run, path, producer_build=BUILD)
    return path


def _document() -> dict:
    """A minimal, synthetic SF Analyzer audit document with hostile text."""
    return {
        "schema_version": "2.0",
        "tool": {"name": "seohead-sf-analyzer", "version": "test"},
        "run": {
            "input_mode": "parse-exports",
            "generated_at": "2026-10-04T00:00:00Z",
            "crawl_partial": False,
            "exports_missing": ["images"],
            "checks_skipped": [{"id": "IMG_MISSING_ALT", "reason": "missing export: images"}],
            "checks_disabled": [{"id": "DESC_TOO_LONG", "reason": "disabled by profile"}],
        },
        "summary": {
            "totals": {"urls_crawled": 2},
            "check_coverage": {
                "checks_total": 162,
                "checks_fired": 1,
                "checks_skipped": 1,
                "checks_disabled": 1,
                "checks_disabled_ids": ["DESC_TOO_LONG"],
                "checks_silent": 1,
                "checks_silent_ids": ["H1_MISSING"],
                "coverage": 0.988,
            },
        },
        "issues": [
            {
                "id": "ISSUE-000001",
                "check": "TITLE_MISSING",
                "severity": "warning",
                "source": "SF:internal_all",
                "message": 'Quotes "x" and <angles> & ampersands — café 日本語',
                "target_url": "https://example.com/a",
                "status_code": 200,
                "occurrences_count": 1,
                "locations": [{"source_url": "https://example.com/"}],
                "details": {},
                "fix_hint": "=cmd|'/c calc'!A1",
            }
        ],
        "pages": [
            {
                "url": "https://example.com/a?x=1&y=<2>",
                "status_code": 200,
                "indexability": "Indexable",
                "indexability_status": "",
                "content_type": "text/html; charset=utf-8",
                "metrics": {"title": "T <a> & b"},
                "issues": [],
                "issue_ids": [],
            },
            {"url": "https://example.com/b", "status_code": 404},
        ],
        "groups": [],
    }


def test_json_envelope_records_and_provenance(artifact, legacy_run, tmp_path):
    audit = json.loads((legacy_run / "audit.json").read_text())
    out = tmp_path / "export.json"
    result = export_scan_data(str(artifact), str(out), fmt="json")
    assert result["ok"], result
    document = json.loads(out.read_text(encoding="utf-8"))
    assert document["format"] == EXPORT_FORMAT_VERSION
    provenance = document["provenance"]
    assert provenance["input_kind"] == "scan.v1"
    assert provenance["scan_uuid"]
    assert provenance["audit"]["state"] == "present"
    records = document["records"]
    assert list(records) == ["pages", "links", "findings"]
    assert [row["url"] for row in records["pages"]] == [page["url"] for page in audit["pages"]]
    assert len(records["links"]) == 3
    assert len(records["findings"]) == len(audit["issues"])
    assert document["statistics"]["rows"] == {
        "pages": len(audit["pages"]),
        "links": 3,
        "findings": len(audit["issues"]),
    }
    coverage = document["coverage"]
    assert coverage["records"]["pages"]["exported"] is True
    assert coverage["records"]["pages"]["source_rows"] == len(audit["pages"])
    checks = coverage["checks"]
    assert checks["state"] in {"complete", "partial"}
    assert checks["ran"] or checks["skipped"] or checks["disabled"]
    for entry in checks["skipped"]:
        assert set(entry) == {"id", "reason"}


def test_selected_fields_identical_across_formats(artifact, tmp_path):
    fields = {"pages": ["url", "title"], "findings": ["id", "severity"]}
    outputs = {}
    for fmt, name in (
        ("json", "e.json"),
        ("xml", "e.xml"),
        ("csv", "e.csv"),
        ("xlsx", "e.xlsx"),
    ):
        out = tmp_path / name
        result = export_scan_data(
            str(artifact), str(out), fmt=fmt, records=["pages", "findings"], fields=fields
        )
        assert result["ok"], (fmt, result)
        outputs[fmt] = out
    document = json.loads(outputs["json"].read_text(encoding="utf-8"))
    assert all(set(row) == {"url", "title"} for row in document["records"]["pages"])
    assert all(set(row) == {"id", "severity"} for row in document["records"]["findings"])
    tree = ET.parse(outputs["xml"])
    for page in tree.getroot().findall(f".//{{{XML_NAMESPACE}}}page"):
        assert {child.tag.split("}")[-1] for child in page} == {"url", "title"}
    rows = list(
        csv.reader(
            io.StringIO(outputs["csv"].with_name("e.pages.csv").read_text(encoding="utf-8-sig")),
            delimiter=";",
        )
    )
    assert rows[0] == ["url", "title"]
    assert [row[0] for row in rows[1:]] == [r["url"] for r in document["records"]["pages"]]
    from openpyxl import load_workbook

    workbook = load_workbook(outputs["xlsx"], read_only=True)
    pages = workbook["Pages"]
    sheet_rows = list(pages.iter_rows(values_only=True))
    assert list(sheet_rows[0]) == ["url", "title"]
    assert [row[0] for row in sheet_rows[1:]] == [r["url"] for r in document["records"]["pages"]]
    assert "Links" not in workbook.sheetnames


def test_unknown_record_type_fails_before_output(artifact, tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(str(artifact), str(out), fmt="json", records=["bogus"])
    assert result["ok"] is False
    assert "bogus" in result["error"]
    assert not out.exists()
    assert not list(tmp_path.glob("*.tmp")) + list(tmp_path.glob(".*.tmp"))


def test_unknown_field_fails_before_output(artifact, tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(
        str(artifact), str(out), fmt="json", fields={"pages": ["url", "no_such_field"]}
    )
    assert result["ok"] is False
    assert "no_such_field" in result["error"]
    assert "pages" in result["error"]
    assert not out.exists()
    result = export_scan_data(str(artifact), str(out), fmt="json", fields={"widgets": ["url"]})
    assert result["ok"] is False
    assert "widgets" in result["error"]
    assert not out.exists()


def test_field_known_for_another_type_is_distinguished_from_unknown(artifact, tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(
        artifact, out, fmt="json", records=["links"], fields={"links": ["title", "invented"]}
    )
    assert result["ok"] is False
    assert "invalid for 'links': title" in result["error"]
    assert "unknown field(s): invented" in result["error"]
    assert not out.exists()


def test_fields_for_unselected_record_type_fails(artifact, tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(
        str(artifact),
        str(out),
        fmt="json",
        records=["pages"],
        fields={"links": ["source"]},
    )
    assert result["ok"] is False
    assert "links" in result["error"]
    assert not out.exists()


def test_document_input_has_no_links_and_preserves_check_states(tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(_document(), str(out), fmt="json")
    assert result["ok"], result
    document = json.loads(out.read_text(encoding="utf-8"))
    assert list(document["records"]) == ["pages", "findings"]
    records = document["coverage"]["records"]
    assert records["links"]["state"] == "unavailable"
    assert records["links"]["exported"] is False
    assert records["links"]["source_rows"] is None
    checks = document["coverage"]["checks"]
    assert "TITLE_MISSING" in checks["ran"]
    assert "H1_MISSING" in checks["ran"]
    assert checks["skipped"] == [{"id": "IMG_MISSING_ALT", "reason": "missing export: images"}]
    assert checks["disabled"] == [{"id": "DESC_TOO_LONG", "reason": "disabled by profile"}]
    assert document["coverage"]["limitations"] == [{"kind": "missing_export", "name": "images"}]
    assert document["provenance"]["input_kind"] == "audit_document"


def test_explicit_unavailable_record_type_fails(tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(_document(), str(out), fmt="json", records=["links"])
    assert result["ok"] is False
    assert "links" in result["error"]
    assert not out.exists()


def test_absent_field_stays_distinct_from_measured_zero(artifact, tmp_path):
    out = tmp_path / "e.json"
    result = export_scan_data(
        str(artifact),
        str(out),
        fmt="json",
        records=["pages"],
        fields={"pages": ["url", "crawl_depth", "canonical_outside_head"]},
    )
    assert result["ok"], result
    rows = json.loads(out.read_text(encoding="utf-8"))["records"]["pages"]
    assert rows[0]["crawl_depth"] == 0
    assert "canonical_outside_head" in rows[0]
    assert rows[0]["canonical_outside_head"] is None


def test_field_selection_cannot_remove_provenance_or_coverage(artifact, tmp_path):
    out = tmp_path / "e.json"
    export_scan_data(
        str(artifact),
        str(out),
        fmt="json",
        records=["pages"],
        fields={"pages": ["url"]},
    )
    document = json.loads(out.read_text(encoding="utf-8"))
    assert document["format"] == EXPORT_FORMAT_VERSION
    assert document["provenance"]["scan_uuid"]
    assert document["coverage"]["checks"]["state"] in {"complete", "partial"}
    assert document["coverage"]["records"]["links"]["exported"] is False
    assert document["projection"]["pages"] == ["url"]
    assert all(set(row) == {"url"} for row in document["records"]["pages"])


def test_xml_versioned_shape_and_exact_text_round_trip(tmp_path):
    out = tmp_path / "e.xml"
    result = export_scan_data(_document(), str(out), fmt="xml")
    assert result["ok"], result
    raw = out.read_bytes()
    assert raw.startswith(b'<?xml version="1.0" encoding="UTF-8"?>')
    assert b"<!DOCTYPE" not in raw and b"<!ENTITY" not in raw
    root = ET.parse(out).getroot()
    assert root.tag == f"{{{XML_NAMESPACE}}}scan-export"
    assert root.attrib["format"] == EXPORT_FORMAT_VERSION
    ns = f"{{{XML_NAMESPACE}}}"
    assert [child.tag for child in root] == [
        f"{ns}provenance",
        f"{ns}projection",
        f"{ns}statistics",
        f"{ns}coverage",
        f"{ns}records",
    ]
    pages = root.findall(f".//{ns}records/{ns}pages/{ns}page")
    assert len(pages) == 2
    assert pages[0].find(f"{ns}url").text == "https://example.com/a?x=1&y=<2>"
    absent = pages[1].find(f"{ns}indexability")
    assert absent is not None and absent.get("state") == "absent" and absent.text is None
    finding = root.find(f".//{ns}records/{ns}findings/{ns}finding")
    assert finding.find(f"{ns}message").text == 'Quotes "x" and <angles> & ampersands — café 日本語'
    assert finding.find(f"{ns}fix_hint").text == "=cmd|'/c calc'!A1"
    locations = finding.find(f"{ns}locations")
    assert locations is not None and locations.get("format") == "json"
    assert json.loads(locations.text) == [{"source_url": "https://example.com/"}]


def test_xml_scalars_keep_explicit_types_and_zero(tmp_path):
    out = tmp_path / "typed.xml"
    result = export_scan_data(_document(), out, fmt="xml")
    assert result["ok"], result
    root = ET.parse(out).getroot()
    ns = f"{{{XML_NAMESPACE}}}"
    page = root.find(f".//{ns}records/{ns}pages/{ns}page")
    assert page.find(f"{ns}url").attrib == {"type": "string"}
    assert page.find(f"{ns}status_code").attrib == {"type": "integer"}
    assert page.find(f"{ns}status_code").text == "200"
    assert root.find(f".//{ns}provenance/{ns}run/{ns}crawl_partial").attrib == {"type": "boolean"}
    assert root.find(f".//{ns}coverage/{ns}checks/{ns}coverage").attrib == {"type": "number"}


def test_xml_escaping_does_not_mutate_values(tmp_path):
    out = tmp_path / "e.xml"
    export_scan_data(_document(), str(out), fmt="xml")
    raw = out.read_text(encoding="utf-8")
    assert "&amp;y=&lt;2&gt;" in raw
    assert "café 日本語" in raw


def test_xml_round_trips_carriage_return_without_normalizing_to_lf(tmp_path):
    document = _document()
    document["issues"][0]["message"] = "before\rafter\r\nnext\tline 🙂 e\u0301"
    out = tmp_path / "cr.xml"
    result = export_scan_data(document, out, fmt="xml")
    assert result["ok"], result
    root = ET.parse(out).getroot()
    ns = f"{{{XML_NAMESPACE}}}"
    value = root.find(f".//{ns}records/{ns}findings/{ns}finding/{ns}message")
    assert value is not None and value.text == document["issues"][0]["message"]
    assert b"&#13;" in out.read_bytes()


def test_xml_forbidden_control_refuses_before_publishing(tmp_path):
    document = _document()
    document["issues"][0]["message"] = "before\x01after"
    out = tmp_path / "control.xml"
    result = export_scan_data(document, out, fmt="xml")
    assert result["ok"] is False
    assert "findings[0].message" in result["error"]
    assert "U+0001" in result["error"]
    assert not out.exists()
    assert not list(tmp_path.glob(".*.tmp"))


def test_xml_arbitrary_metadata_keys_are_well_formed_and_reversible(tmp_path):
    document = _document()
    document["summary"]["totals"]["bad <key>&"] = 0
    out = tmp_path / "keys.xml"
    result = export_scan_data(document, out, fmt="xml")
    assert result["ok"], result
    root = ET.parse(out).getroot()
    ns = f"{{{XML_NAMESPACE}}}"
    entry = root.find(f".//{ns}statistics/{ns}run/{ns}totals/{ns}entry")
    assert entry is not None
    assert entry.attrib == {"type": "integer", "key": "bad <key>&"}
    assert entry.text == "0"


def test_xml_invalid_metadata_key_refuses_without_publishing(tmp_path):
    document = _document()
    document["summary"]["totals"]["bad\x01key"] = 1
    out = tmp_path / "invalid-key.xml"
    result = export_scan_data(document, out, fmt="xml")
    assert result["ok"] is False
    assert "U+0001" in result["error"]
    assert not out.exists()


def test_csv_files_and_manifest(artifact, tmp_path):
    out = tmp_path / "export.csv"
    result = export_scan_data(str(artifact), str(out), fmt="csv")
    assert result["ok"], result
    names = {path.name for path in tmp_path.iterdir() if path.name != "scan.sqlite"}
    assert {
        "export.pages.csv",
        "export.links.csv",
        "export.findings.csv",
        "export.manifest.csv",
    } <= names
    manifest = (tmp_path / "export.manifest.csv").read_text(encoding="utf-8-sig")
    assert f"format;{EXPORT_FORMAT_VERSION}" in manifest
    assert "file.pages;export.pages.csv" in manifest
    assert "statistics.rows.pages;2" in manifest
    assert "coverage.records.pages.state" in manifest
    rows = (tmp_path / "export.pages.csv").read_text(encoding="utf-8-sig").splitlines()
    assert rows[0].startswith("url;")
    assert len(rows) == 3


def test_csv_formula_leading_values_are_neutralized(tmp_path):
    out = tmp_path / "e.csv"
    result = export_scan_data(_document(), str(out), fmt="csv", records=["findings"])
    assert result["ok"], result
    text = (tmp_path / "e.findings.csv").read_text(encoding="utf-8-sig")
    assert "\\F=cmd|'/c calc'!A1" in text
    assert "\n=cmd" not in text


def test_xlsx_bounded_workbook_and_formula_neutralization(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "e.xlsx"
    result = export_scan_data(_document(), str(out), fmt="xlsx")
    assert result["ok"], result
    workbook = load_workbook(out, read_only=True)
    assert set(workbook.sheetnames) == {"Summary", "Pages", "Findings"}
    findings = list(workbook["Findings"].iter_rows(values_only=True))
    assert list(findings[0]) == list(export_module.FINDING_FIELDS)
    hint = findings[1][list(export_module.FINDING_FIELDS).index("fix_hint")]
    assert hint == "\\F=cmd|'/c calc'!A1"
    summary = {
        row[0]: row[1]
        for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)
        if len(row) >= 2
    }
    assert summary["format"] == EXPORT_FORMAT_VERSION
    assert summary["statistics.rows.findings"] == 1


def test_tabular_null_empty_formula_and_escape_values_are_distinct(tmp_path):
    from openpyxl import load_workbook

    values = [None, "", r"\N", r"\E", r"\F=literal", "=SUM(1,1)", "\t=SUM(1,1)", "plain"]
    expected = [
        r"\N",
        r"\E",
        r"\\N",
        r"\\E",
        r"\\F=literal",
        r"\F=SUM(1,1)",
        "\\F\t=SUM(1,1)",
        "plain",
    ]
    document = _document()
    document["pages"] = [
        {"url": f"https://example.com/{index}", "indexability_status": value}
        for index, value in enumerate(values)
    ]
    fields = {"pages": ["url", "indexability_status"]}
    csv_out = tmp_path / "values.csv"
    xlsx_out = tmp_path / "values.xlsx"
    for fmt, out in (("csv", csv_out), ("xlsx", xlsx_out)):
        result = export_scan_data(document, out, fmt=fmt, records=["pages"], fields=fields)
        assert result["ok"], result
    with (tmp_path / "values.pages.csv").open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.reader(handle, delimiter=";"))
    assert [row[1] for row in csv_rows[1:]] == expected
    workbook = load_workbook(xlsx_out, read_only=True)
    xlsx_rows = list(workbook["Pages"].iter_rows(values_only=True))
    assert [row[1] for row in xlsx_rows[1:]] == expected
    assert all(not value.startswith("=") for value in expected)


def test_xlsx_row_limit_fails_before_writing(artifact, tmp_path, monkeypatch):
    monkeypatch.setattr(export_module, "EXCEL_MAX_ROWS", 3)
    out = tmp_path / "e.xlsx"
    result = export_scan_data(str(artifact), str(out), fmt="xlsx")
    assert result["ok"] is False
    assert "row limit" in result["error"]
    assert not out.exists()
    assert not list(tmp_path.glob(".*.tmp"))


def test_xlsx_cell_text_limit_fails_cleanly(tmp_path):
    document = _document()
    document["pages"][0]["content_type"] = "x" * 40_000
    out = tmp_path / "e.xlsx"
    result = export_scan_data(document, str(out), fmt="xlsx")
    assert result["ok"] is False
    assert "cell-text limit" in result["error"]
    assert not out.exists()
    assert not list(tmp_path.glob(".*.tmp"))


def test_serialization_failure_leaves_no_final_artifact(tmp_path):
    document = _document()
    document["pages"][0]["metrics"] = {"unserializable": object()}
    out = tmp_path / "e.json"
    result = export_scan_data(document, str(out), fmt="json")
    assert result["ok"] is False
    assert not out.exists()
    assert not list(tmp_path.glob(".*.tmp"))
    out_xml = tmp_path / "e.xml"
    result = export_scan_data(document, str(out_xml), fmt="xml")
    assert result["ok"] is False
    assert not out_xml.exists()
    assert not list(tmp_path.glob(".*.tmp"))


def test_existing_destination_is_refused(artifact, tmp_path):
    out = tmp_path / "e.json"
    out.write_text("already here")
    result = export_scan_data(str(artifact), str(out), fmt="json")
    assert result["ok"] is False
    assert "already exists" in result["error"]
    assert out.read_text() == "already here"


def test_symlink_input_alias_and_csv_child_conflicts_preserve_bytes(artifact, tmp_path):
    retained = artifact.read_bytes()
    result = export_scan_data(artifact, artifact, fmt="json")
    assert result["ok"] is False and artifact.read_bytes() == retained

    outside = tmp_path / "untouched.txt"
    outside.write_bytes(b"keep")
    output_link = tmp_path / "linked.xml"
    output_link.symlink_to(outside)
    result = export_scan_data(_document(), output_link, fmt="xml")
    assert result["ok"] is False and outside.read_bytes() == b"keep"

    csv_child = tmp_path / "batch.pages.csv"
    csv_child.symlink_to(outside)
    result = export_scan_data(_document(), tmp_path / "batch.csv", fmt="csv", records=["pages"])
    assert result["ok"] is False and outside.read_bytes() == b"keep"
    assert not (tmp_path / "batch.manifest.csv").exists()

    result = export_scan_data(_document(), tmp_path / "other.csv", fmt="csv", records=["../pages"])
    assert result["ok"] is False
    assert not list(tmp_path.glob("other*"))


def test_json_chunks_consume_records_incrementally():
    produced = {"count": 0}

    def rows():
        for index in range(100_000):
            produced["count"] += 1
            yield {"url": f"https://example.com/{index}"}

    head = {
        "format": EXPORT_FORMAT_VERSION,
        "provenance": {"input_kind": "synthetic"},
        "projection": {"pages": ["url"]},
        "statistics": {"record_types": ["pages"], "rows": {"pages": 100_000}, "run": {}},
        "coverage": {"records": {}, "checks": {}},
    }
    chunks = export_module._json_chunks(head, [("pages", ("url",), rows())])
    record_chunks = 0
    for chunk in chunks:
        if b"example.com" in chunk:
            record_chunks += 1
            if record_chunks == 3:
                break
    assert record_chunks == 3
    assert produced["count"] == 3


def test_deterministic_order_and_repeated_export(artifact, tmp_path):
    first = tmp_path / "first.json"
    second = tmp_path / "second.json"
    assert export_scan_data(str(artifact), str(first), fmt="json")["ok"]
    assert export_scan_data(str(artifact), str(second), fmt="json")["ok"]
    assert first.read_bytes() == second.read_bytes()


def test_site_audit_document_is_refused(tmp_path):
    document = {"schema": "seohead.site-audit/1", "findings": [], "pages": [], "summary": {}}
    out = tmp_path / "e.json"
    result = export_scan_data(document, str(out), fmt="json")
    assert result["ok"] is False
    assert "site-audit" in result["error"]
    assert not out.exists()


def test_scan_export_handler_end_to_end(artifact, tmp_path):
    from seohead.servers import handlers

    out = tmp_path / "e.json"
    result = handlers.scan_export(
        input_path=str(artifact), out=str(out), format="json", records="pages,links"
    )
    assert result["ok"], result
    document = json.loads(out.read_text(encoding="utf-8"))
    assert list(document["records"]) == ["pages", "links"]


def test_cli_scan_export_executes(artifact, tmp_path, monkeypatch, capsys):
    from seohead import cli

    monkeypatch.setenv("SEOHEAD_RUN_LOG", "off")
    out = tmp_path / "e.json"
    rc = cli.main(["scan-export", "--scan", str(artifact), "--out", str(out), "--format", "json"])
    assert rc == 0
    envelope = json.loads(capsys.readouterr().out)
    assert envelope["ok"] is True
    assert json.loads(out.read_text(encoding="utf-8"))["format"] == EXPORT_FORMAT_VERSION
    out2 = tmp_path / "e2.xml"
    rc = cli.main(
        ["scan", "export", "--scan", str(artifact), "--out", str(out2), "--format", "xml"]
    )
    assert rc == 0
    assert out2.exists()
