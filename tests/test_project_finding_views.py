"""Offline persistence, deterministic projections, and report reuse for saved finding views."""

from __future__ import annotations

import copy
import json

import pytest

from seohead.audit.site import SCHEMA
from seohead.projects.finding_views import list_views, save_view, show_view
from seohead.projects.workspace import create_project

SITE = "https://example.test/"


def definition(name="triage", **values):
    return {
        "name": name,
        "filters": {},
        "sort": {"field": "severity", "direction": "desc"},
        "columns": ["severity", "check", "url", "text"],
        "page_size": 2,
        **values,
    }


def audit_document():
    return {
        "schema": SCHEMA,
        "domain": "example.test",
        "url": SITE,
        "generated_at": "2026-10-03T12:00:00Z",
        "run": {
            "scan_uuid": "scan-synthetic-757",
            "finding_exclusion_policy": [
                {"id": "synthetic-rule", "reason": "synthetic source exclusion"}
            ],
            "crawl_config": {
                "scope.segments": [{"name": "blog", "prefix": "/blog/"}],
                "scope.segments_only": [],
            },
        },
        "findings": [
            {
                "severity": "critical",
                "check": "BROKEN_PAGE_4XX",
                "url": "https://example.test/blog/a",
                "text": "Synthetic broken URL",
                "status_code": 404,
                "occurrences_count": 1,
                "fix_hint": "Repair the fixture link",
            },
            {
                "severity": "warning",
                "check": "MISSING_TITLE",
                "url": "https://example.test/blog/b",
                "text": "Synthetic missing title",
                "status_code": 200,
                "occurrences_count": 1,
                "fix_hint": "Add a title",
            },
            {
                "severity": "warning",
                "check": "MISSING_TITLE",
                "url": "https://example.test/blog/c",
                "text": "Synthetic second missing title",
                "status_code": 200,
                "occurrences_count": 1,
                "fix_hint": "Add a title",
            },
            {
                "severity": "warning",
                "check": "MISSING_TITLE",
                "text": "Audit-wide fixture without a target URL",
            },
        ],
        "pages": [
            {"url": "https://example.test/blog/a", "status": 404, "title": "A"},
            {"url": "https://example.test/blog/b", "status": 200, "title": "B"},
            {"url": "https://example.test/blog/c", "status": 200, "title": "C"},
        ],
        "summary": {
            "pages_checked": 3,
            "findings_total": 4,
            "findings_by_severity": {"critical": 1, "warning": 3, "notice": 0},
            "finding_exclusions": {
                "rules_configured": 1,
                "suppressed_total": 2,
                "by_rule": {"synthetic-rule": 2},
                "by_check": {"SYNTHETIC_SUPPRESSED": 2},
                "by_severity": {"notice": 2},
            },
        },
    }


@pytest.fixture
def project(tmp_path):
    root = tmp_path / "project"
    create_project(root, SITE)
    return root


def test_view_round_trip_identity_revisions_and_expected_revision(project):
    saved = save_view(project, definition(), expected_revision=0)
    view_id = saved["view"]["id"]
    assert saved["config_revision"] == 1
    assert saved["view"]["revision"] == 1

    # A new process/re-read sees the same stable view identity and definition.
    assert show_view(project, "triage")["view"]["id"] == view_id
    assert list_views(project)["views"][0]["definition"] == definition()

    updated = save_view(
        project,
        definition(columns=["check", "url"], page_size=10),
        expected_revision=1,
    )
    assert updated["view"]["id"] == view_id
    assert updated["view"]["revision"] == 2
    assert updated["config_revision"] == 2

    before = (project / "finding-views.json").read_bytes()
    with pytest.raises(ValueError, match="revision conflict"):
        save_view(project, definition(name="second"), expected_revision=1)
    assert (project / "finding-views.json").read_bytes() == before


@pytest.mark.parametrize(
    "bad",
    [
        {"name": "x", "filters": {"expression": "1=1"}},
        {"name": "x", "filters": {"url_regex": [".*"]}},
        {"name": "x", "sort": {"field": "__import__", "direction": "asc"}},
        {"name": "x", "sort": {"field": ["url"], "direction": "asc"}},
        {"name": "x", "columns": ["severity", "sql"]},
        {"name": "x", "filters": {"severity": []}},
        {"name": "x", "page_size": 1001},
    ],
)
def test_closed_view_schema_rejects_expressions_and_unsupported_values(project, bad):
    before = (
        (project / "finding-views.json").read_bytes()
        if (project / "finding-views.json").exists()
        else None
    )
    with pytest.raises(ValueError):
        save_view(project, bad, expected_revision=0)
    after = (
        (project / "finding-views.json").read_bytes()
        if (project / "finding-views.json").exists()
        else None
    )
    assert after == before


def test_filter_sort_projection_and_pagination_are_repeatable(project):
    save_view(
        project,
        definition(
            filters={
                "severity": ["warning"],
                "check": ["MISSING_TITLE"],
                "url": ["https://example.test/blog/b", "https://example.test/blog/c"],
                "segment": ["blog"],
            },
            sort={"field": "url", "direction": "desc"},
            columns=["check", "url", "segment", "text"],
            page_size=1,
        ),
        expected_revision=0,
    )
    from seohead.projects.finding_views import apply_view_to_audit

    document = audit_document()
    original = copy.deepcopy(document)
    original = copy.deepcopy(document)
    first = apply_view_to_audit(project, "triage", document)
    first_again = apply_view_to_audit(project, "triage", document)
    next_page = apply_view_to_audit(project, "triage", document, offset=1)

    assert document == original
    assert first["state"] == "partial"
    assert first["counts"] == {
        "source": 4,
        "matched": 2,
        "returned": 1,
        "filtered": 2,
        "missing_filter_fields": {"url": 1},
        "missing_projection_fields": {},
    }
    assert first["pagination"] == {
        "offset": 0,
        "page_size": 1,
        "has_more": True,
        "truncated": True,
        "next_offset": 1,
    }
    assert first["items"] == first_again["items"]
    assert first["items"][0]["fields"] == {
        "check": "MISSING_TITLE",
        "url": "https://example.test/blog/c",
        "segment": "blog",
        "text": "Synthetic second missing title",
    }
    assert next_page["items"][0]["fields"]["url"] == "https://example.test/blog/b"
    assert next_page["items"][0]["finding_id"] != first["items"][0]["finding_id"]


def test_severity_sort_directions_follow_severity_strength(project):
    from seohead.projects.finding_views import apply_view_to_audit

    document = audit_document()
    document["findings"].append(
        {
            "severity": "notice",
            "check": "OPTIONAL_FIXTURE",
            "url": "https://example.test/blog/notice",
            "text": "Synthetic notice",
        }
    )
    save_view(
        project,
        definition(columns=["severity"], page_size=10),
        expected_revision=0,
    )

    descending = apply_view_to_audit(project, "triage", document)
    assert [item["fields"]["severity"] for item in descending["items"]] == [
        "critical",
        "warning",
        "warning",
        "warning",
        "notice",
    ]

    save_view(
        project,
        definition(
            columns=["severity"], page_size=10, sort={"field": "severity", "direction": "asc"}
        ),
        expected_revision=1,
    )
    ascending = apply_view_to_audit(project, "triage", document)
    assert [item["fields"]["severity"] for item in ascending["items"]] == [
        "notice",
        "warning",
        "warning",
        "warning",
        "critical",
    ]


def test_missing_filter_fields_are_counted_and_never_match(project):
    save_view(
        project,
        definition(
            filters={"url": ["https://example.test/blog/a"]},
            columns=["severity", "segment"],
        ),
        expected_revision=0,
    )
    from seohead.projects.finding_views import apply_view_to_audit

    no_segments = {**audit_document(), "run": {"scan_uuid": "scan-no-segments"}}
    result = apply_view_to_audit(project, "triage", no_segments)
    assert result["state"] == "partial"
    assert result["counts"]["matched"] == 1
    assert result["counts"]["missing_filter_fields"] == {"url": 1}
    assert result["counts"]["missing_projection_fields"] == {"segment": 1}
    assert result["items"][0]["missing_fields"] == ["segment"]


def test_segment_filter_requires_and_uses_source_segment_definitions(project):
    from seohead.projects.finding_views import apply_view_to_audit

    save_view(
        project,
        definition(filters={"segment": ["blog"]}, columns=["url", "segment"]),
        expected_revision=0,
    )
    result = apply_view_to_audit(project, "triage", audit_document())
    assert result["counts"]["matched"] == 3
    assert all(item["fields"]["segment"] == "blog" for item in result["items"])

    no_segments = {**audit_document(), "run": {"scan_uuid": "scan-no-segments"}}
    with pytest.raises(ValueError, match="segment filtering is unavailable"):
        apply_view_to_audit(project, "triage", no_segments)


def test_segment_view_uses_the_audits_declared_segment_rules(project):
    from seohead.projects.finding_views import apply_view_to_audit

    save_view(
        project,
        definition(filters={"segment": ["blog"]}, columns=["url", "segment"]),
        expected_revision=0,
    )
    document = audit_document()
    document["run"]["crawl_config"]["scope.segments"] = [
        {"name": "blog", "pattern": r"https://example\.test/blog/"}
    ]
    result = apply_view_to_audit(project, "triage", document)
    assert result["counts"]["matched"] == 3
    assert all(item["fields"]["segment"] == "blog" for item in result["items"])


def test_json_and_markdown_reports_apply_the_same_projection_without_mutating_audit(
    project, tmp_path
):
    from seohead.reports import build_report

    save_view(
        project,
        definition(
            name="warnings",
            filters={"severity": ["warning"]},
            sort={"field": "check", "direction": "asc"},
            columns=["check", "url"],
            page_size=10,
        ),
        expected_revision=0,
    )
    document = audit_document()
    original = copy.deepcopy(document)
    markdown_path = tmp_path / "filtered.md"
    md = build_report(
        document, fmt="md", path=str(markdown_path), project=str(project), view="warnings"
    )
    assert md["ok"] is True, md
    assert md["findings"] == 3
    rendered = markdown_path.read_text(encoding="utf-8")
    assert "Filtered finding view" in rendered
    assert "Check | URL" in rendered
    assert "BROKEN_PAGE_4XX" not in rendered
    assert "https://example.test/blog/b" in rendered
    assert "Audit totals, evidence coverage, and scores describe the full source audit." in rendered
    assert "2 excluded findings" in rendered
    assert "synthetic-rule=2" in rendered
    assert document == original


def test_csv_xlsx_and_docx_reports_use_saved_projection_and_state(project, tmp_path):
    import csv

    from docx import Document
    from openpyxl import load_workbook

    from seohead.reports import build_report

    save_view(
        project,
        definition(
            name="warning-urls",
            filters={"severity": ["warning"]},
            sort={"field": "url", "direction": "asc"},
            columns=["check", "url"],
            page_size=1,
        ),
        expected_revision=0,
    )
    document = audit_document()
    original = copy.deepcopy(document)

    csv_path = tmp_path / "filtered.csv"
    csv_result = build_report(
        document, fmt="csv", path=str(csv_path), project=str(project), view="warning-urls"
    )
    assert csv_result["ok"] is True, csv_result
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=";"))
    assert rows[0] == ["Check", "URL"]
    assert len(rows) == 2
    scope = csv_path.with_suffix(".scope.csv").read_text(encoding="utf-8-sig")
    assert "finding view" in scope and "More matching rows remain" in scope
    assert "synthetic-rule=2" in scope

    xlsx_path = tmp_path / "filtered.xlsx"
    xlsx_result = build_report(
        document, fmt="xlsx", path=str(xlsx_path), project=str(project), view="warning-urls"
    )
    assert xlsx_result["ok"] is True, xlsx_result
    workbook = load_workbook(xlsx_path, read_only=True)
    findings = workbook["Findings"]
    assert [cell.value for cell in findings[1]] == ["Check", "URL"]
    assert findings.max_row == 2
    summary = workbook["Summary"]
    assert any(
        "Saved finding view" in str(cell.value) for row in summary.iter_rows() for cell in row
    )
    workbook.close()

    docx_path = tmp_path / "filtered.docx"
    docx_result = build_report(
        document, fmt="docx", path=str(docx_path), project=str(project), view="warning-urls"
    )
    assert docx_result["ok"] is True, docx_result
    word = Document(docx_path)
    body = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "Saved finding view warning-urls" in body
    findings_table = next(
        table
        for table in word.tables
        if [cell.text for cell in table.rows[0].cells] == ["Check", "URL"]
    )
    assert len(findings_table.rows) == 2

    json_path = tmp_path / "filtered.json"
    output = build_report(
        document, fmt="json", path=str(json_path), project=str(project), view="warning-urls"
    )
    assert output["ok"] is True, output
    envelope = json.loads(json_path.read_text(encoding="utf-8"))
    assert envelope["schema"] == "seohead.finding-view/1"
    assert envelope["view"]["name"] == "warning-urls"
    assert envelope["view"]["counts"]["source"] == 4
    assert (
        envelope["source"]["finding_exclusion_policy"]
        == document["run"]["finding_exclusion_policy"]
    )
    assert envelope["source"]["finding_exclusions"] == document["summary"]["finding_exclusions"]
    assert set(envelope["rows"][0]["fields"]) == {"check", "url"}
    assert document == original


def test_cli_and_mcp_expose_the_same_view_contract(monkeypatch, project):
    from seohead import cli
    from seohead.servers import handlers
    from seohead.servers.mcp_server import build_server

    config = json.dumps(definition())
    args = cli.build_parser().parse_args(
        [
            "project",
            "view-save",
            "--directory",
            str(project),
            "--expected-revision",
            "0",
            "--input",
            json.dumps({"view": json.loads(config)}),
        ]
    )
    command = "project-" + args.project_command
    handler_name, kwargs = cli._build_kwargs(command, args)
    assert handler_name == "project_view_save"
    assert kwargs == {"directory": str(project), "expected_revision": 0, "view": definition()}

    apply_args = cli.build_parser().parse_args(
        [
            "findings-view",
            "--directory",
            str(project),
            "--name",
            "triage",
            "--audit",
            "audit.json",
            "--offset",
            "2",
        ]
    )
    handler_name, kwargs = cli._build_kwargs("findings-view", apply_args)
    assert handler_name == "findings_view"
    assert kwargs == {
        "directory": str(project),
        "name": "triage",
        "audit": "audit.json",
        "offset": 2,
    }

    captured = {}

    def capture(**kwargs):
        captured.update(kwargs)
        return {"ok": True}

    monkeypatch.setattr(handlers, "project_view_save", capture)
    manager = build_server()._tool_manager
    manager.get_tool("seo_project_view_save").fn(
        directory=str(project), view=definition(), expected_revision=0
    )
    assert captured == {"directory": str(project), "view": definition(), "expected_revision": 0}

    captured.clear()
    monkeypatch.setattr(handlers, "findings_view", capture)
    manager.get_tool("seo_findings_view").fn(
        directory=str(project), name="triage", audit={"schema": SCHEMA}, offset=2
    )
    assert captured == {
        "directory": str(project),
        "name": "triage",
        "audit": {"schema": SCHEMA},
        "offset": 2,
    }


def test_cli_end_to_end_saves_applies_and_reports_a_view_offline(
    monkeypatch, project, tmp_path, capsys
):
    from seohead import cli

    monkeypatch.setenv("SEOHEAD_RUN_LOG", "off")
    config = definition(
        name="warning-urls",
        filters={"severity": ["warning"]},
        sort={"field": "url", "direction": "asc"},
        columns=["check", "url"],
        page_size=1,
    )
    assert (
        cli.main(
            [
                "project-view-save",
                "--directory",
                str(project),
                "--expected-revision",
                "0",
                "--input",
                json.dumps({"view": config}),
            ]
        )
        == 0
    )
    saved = json.loads(capsys.readouterr().out)
    assert saved["view"]["revision"] == 1

    source = audit_document()
    audit_path = tmp_path / "audit.json"
    audit_path.write_text(json.dumps(source), encoding="utf-8")
    assert (
        cli.main(
            [
                "findings-view",
                "--directory",
                str(project),
                "--name",
                "warning-urls",
                "--audit",
                str(audit_path),
            ]
        )
        == 0
    )
    applied = json.loads(capsys.readouterr().out)
    assert applied["counts"]["matched"] == 3
    assert applied["pagination"]["has_more"] is True

    report = project / "reports" / "view.md"
    assert (
        cli.main(
            [
                "report-build",
                "--audit",
                str(audit_path),
                "--project",
                str(project),
                "--view",
                "warning-urls",
                "--format",
                "md",
                "--out",
                str(report),
            ]
        )
        == 0
    )
    result = json.loads(capsys.readouterr().out)
    assert result["finding_view"]["name"] == "warning-urls"
    assert result["finding_view"]["pagination"]["truncated"] is True
    assert "Saved finding view" in report.read_text(encoding="utf-8")
    assert json.loads(audit_path.read_text(encoding="utf-8")) == source


def test_findings_view_reads_a_synthetic_validated_scan_offline(project):
    from seohead.servers.handlers import findings_view
    from seohead.storage import read_audit
    from tests.test_scan_reanalysis_integration import _source

    save_view(project, definition(columns=["check", "url"]), expected_revision=0)
    scan = project / "scans" / "synthetic.sqlite"
    _source(scan)
    result = findings_view(str(project), "triage", str(scan))
    assert result["ok"] is True
    assert result["source"]["rows_key"] == "issues"
    assert result["source"]["identity"]
    assert result["counts"]["source"] == len(read_audit(scan)["issues"])
