"""Synthetic, offline tests for the ordered audit.v2 companion."""

from __future__ import annotations

import json
import zipfile
from contextlib import closing

import pytest

from seohead.storage import open_scan
from seohead.storage.audit_v2 import (
    AuditV2Error,
    AuditV2Reader,
    audit_v2_path,
    write_audit_v2,
)
from seohead.storage.native_scan import NativeScan
from tests.test_scan_native import _metadata


def _scan(path):
    with NativeScan.create(path, **_metadata()) as scan:
        assert scan.finish_without_audit("audit.v2 test")
    with closing(open_scan(path, require_audit=False)) as con:
        row = con.execute("SELECT * FROM scan WHERE singleton=1").fetchone()
    return {
        "scan_uuid": row["scan_uuid"],
        "evidence_revision": row["evidence_revision"],
        "analyzer_version": row["writer_version"],
        "analyzer_revision": row["writer_revision"],
    }


def test_audit_v2_reopens_ordered_collections_and_complete_document(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    header = {
        "schema_version": "2.0",
        "run": {"source": "https://example.test/", "crawl_partial": False},
        "issues": [],
        "pages": [],
        "groups": [],
        "summary": {"sitemap": {"linked_not_in_sitemap": []}},
    }
    issues = (
        {"check": "TEST", "target_url": f"https://example.test/{i}", "n": i} for i in range(10000)
    )
    pages = [{"url": f"https://example.test/{i}"} for i in range(3)]
    sitemap = ({"url": "https://example.test/orphan"},)
    groups = [{"check": "TEST", "count": 1}]

    written = write_audit_v2(
        scan,
        header,
        {
            "/issues": issues,
            "/pages": pages,
            "/summary/sitemap/linked_not_in_sitemap": sitemap,
            "/groups": groups,
        },
        binding,
    )
    assert written == audit_v2_path(scan)

    with AuditV2Reader(scan) as reader:
        assert reader.count("/issues") == 10000
        assert reader.count("/pages") == 3
        assert reader.count("/groups") == 1
        assert list(reader.iter_collection("/pages")) == list(pages)
        document = reader.materialize_legacy()
        assert len(document["issues"]) == 10000
        assert document["issues"][0]["target_url"].endswith("/0")
        assert document["issues"][-1]["n"] == 9999
        assert document["summary"]["sitemap"]["linked_not_in_sitemap"] == list(sitemap)
        assert document["groups"] == groups
        encoded = "".join(reader.document_chunks())
        assert json.loads(encoded) == document


def test_audit_v2_enforces_legacy_export_limit_while_streaming(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    write_audit_v2(
        scan,
        {"issues": [], "pages": []},
        {"/issues": ({"message": "x" * 1024} for _ in range(100))},
        binding,
    )
    with AuditV2Reader(scan) as reader:
        chunks = reader.document_chunks(max_bytes=4096)
        with pytest.raises(AuditV2Error, match="legacy JSON export exceeds"):
            for _ in chunks:
                pass
        with pytest.raises(AuditV2Error, match="legacy JSON export exceeds"):
            reader.materialize_legacy(max_bytes=4096)


def test_audit_v2_refuses_wrong_scan_binding_and_tampering(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    with pytest.raises(AuditV2Error, match="binding disagrees"):
        write_audit_v2(
            scan,
            {"issues": []},
            {"/issues": []},
            {**binding, "scan_uuid": "00000000-0000-0000-0000-000000000000"},
        )

    write_audit_v2(scan, {"issues": []}, {"/issues": [{"check": "X"}]}, binding)
    import sqlite3

    with sqlite3.connect(audit_v2_path(scan)) as con:
        con.execute("UPDATE items SET value_json='{}'")
    with pytest.raises(AuditV2Error, match="hash"):
        AuditV2Reader(scan)


def test_report_build_streams_all_large_audit_rows(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    issues = [
        {
            "check": "TITLE_MISSING",
            "severity": "warning",
            "target_url": f"https://example.test/{index}",
            "message": "Title is missing",
            "details": {"index": index},
        }
        for index in range(10000)
    ]
    pages = [
        {
            "url": f"https://example.test/{index}",
            "status_code": 200,
            "metrics": {"title": "", "word_count": 2},
        }
        for index in range(10000)
    ]
    header = {
        "schema_version": "2.0",
        "tool": {"version": binding["analyzer_version"]},
        "run": {"source": "https://example.test/", "generated_at": "2026-10-03T00:00:00Z"},
        "summary": {
            "totals": {"urls_crawled": 10000, "issues_total": 10000},
            "by_severity": {"critical": 0, "warning": 10000, "notice": 0},
        },
        "issues": [],
        "pages": [],
        "groups": [],
    }
    write_audit_v2(
        scan,
        header,
        {"/issues": issues, "/pages": pages, "/groups": [{"check": "TITLE_MISSING"}]},
        binding,
    )

    from seohead.reports import build_report

    json_path = tmp_path / "large.json"
    result = build_report(scan, fmt="json", path=str(json_path))
    assert result["ok"] is True
    assert result["findings"] == result["pages"] == 10000
    with json_path.open(encoding="utf-8") as stream:
        exported = json.load(stream)
    assert len(exported["issues"]) == len(exported["pages"]) == 10000
    assert exported["groups"] == [{"check": "TITLE_MISSING"}]
    assert exported["issues"][-1]["target_url"].endswith("/9999")

    csv_path = tmp_path / "large.csv"
    result = build_report(scan, fmt="csv", path=str(csv_path))
    assert result["ok"] is True, result
    assert result["findings"] == result["pages"] == 10000
    assert sum(1 for _ in csv_path.open(encoding="utf-8-sig")) == 10001
    assert sum(1 for _ in csv_path.with_suffix(".pages.csv").open(encoding="utf-8-sig")) == 10001

    xlsx_path = tmp_path / "large.xlsx"
    result = build_report(scan, fmt="xlsx", path=str(xlsx_path))
    assert result["ok"] is True
    with zipfile.ZipFile(xlsx_path) as archive:
        assert "xl/charts/chart1.xml" in archive.namelist()
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True)
    try:
        findings_rows = list(workbook["Findings"].iter_rows(values_only=True))
        pages_rows = list(workbook["Pages"].iter_rows(values_only=True))
        assert len(findings_rows) == len(pages_rows) == 10001
        assert findings_rows[-1][1] == "https://example.test/9999"
        assert pages_rows[-1][0] == "https://example.test/9999"
    finally:
        workbook.close()

    md_path = tmp_path / "large.md"
    result = build_report(scan, fmt="md", path=str(md_path))
    assert result["ok"] is True
    assert "https://example.test/9999" in md_path.read_text(encoding="utf-8")

    docx_path = tmp_path / "large.docx"
    result = build_report(scan, fmt="docx", path=str(docx_path))
    assert result["ok"] is True, result
    from docx import Document

    report = Document(docx_path)
    assert any("Warning — 10000" in paragraph.text for paragraph in report.paragraphs)
    assert any(
        "Showing the first 60 of 10000 pages." in paragraph.text for paragraph in report.paragraphs
    )


def test_json_report_streams_document_larger_than_legacy_ceiling(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    issue_body = "x" * 7000
    header = {
        "schema_version": "2.0",
        "run": {"source": "https://example.test/"},
        "summary": {"totals": {"issues_total": 10000}, "by_severity": {}},
        "issues": [],
        "pages": [],
        "groups": [],
    }
    write_audit_v2(
        scan,
        header,
        {
            "/issues": (
                {
                    "check": "HIGH_VOLUME",
                    "target_url": f"https://example.test/{i}",
                    "message": issue_body,
                }
                for i in range(10000)
            ),
            "/pages": [],
            "/groups": [{"check": "HIGH_VOLUME", "count": 10000}],
        },
        binding,
    )
    assert audit_v2_path(scan).stat().st_size > 64 * 1024 * 1024

    from seohead.reports import build_report

    output = tmp_path / "oversized.json"
    result = build_report(scan, fmt="json", path=str(output))
    assert result["ok"] is True
    assert result["findings"] == 10000
    assert output.stat().st_size > 64 * 1024 * 1024
    with output.open("rb") as stream:
        stream.seek(-2, 2)
        assert stream.read() == b"}}"
    with AuditV2Reader(scan) as reader:
        assert reader.count("/issues") == 10000
        assert list(reader.iter_collection("/issues"))[-1]["target_url"].endswith("/9999")


def test_compare_crawls_accepts_streaming_audit_v2_sources(tmp_path):
    from seohead.servers.handlers import compare_crawls

    before_path = tmp_path / "before.sqlite"
    after_path = tmp_path / "after.sqlite"
    before_binding = _scan(before_path)
    after_binding = _scan(after_path)
    before_header = {"run": {"generated_at": "before"}, "issues": [], "pages": []}
    after_header = {"run": {"generated_at": "after"}, "issues": [], "pages": []}
    before_pages = [
        {"url": "https://example.test/"},
        *({"url": f"https://example.test/{i}"} for i in range(1, 10000)),
    ]
    after_pages = list(before_pages)
    before_issues = (
        {"check": f"CHECK_{i:05d}", "target_url": "https://example.test/"} for i in range(10000)
    )
    after_issues = []
    write_audit_v2(
        before_path,
        before_header,
        {"/issues": before_issues, "/pages": before_pages},
        before_binding,
    )
    write_audit_v2(
        after_path,
        after_header,
        {"/issues": after_issues, "/pages": after_pages},
        after_binding,
    )

    result = compare_crawls(before=str(before_path), after=str(after_path))
    assert result["summary"]["left"] == 10000
    assert len(result["left"]) == 10000
    assert result["summary"]["disappeared"] == 0


def test_audit_v2_saves_and_reopens_populations_above_ten_thousand(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    write_audit_v2(
        scan,
        {"issues": []},
        {
            "/issues": (
                {"check": "HIGH_VOLUME", "target_url": f"https://example.test/{i}"}
                for i in range(50001)
            )
        },
        binding,
    )
    with AuditV2Reader(scan) as reader:
        assert reader.count("/issues") == 50001
        assert sum(1 for _ in reader.iter_collection("/issues")) == 50001


def test_report_refuses_stale_severity_summary_instead_of_claiming_clean(tmp_path):
    scan = tmp_path / "scan.sqlite"
    binding = _scan(scan)
    write_audit_v2(
        scan,
        {
            "schema_version": "2.0",
            "run": {},
            "summary": {
                "totals": {"issues_total": 1, "urls_crawled": 0},
                "by_severity": {"critical": 0, "warning": 0, "notice": 0},
            },
            "issues": [],
            "pages": [],
        },
        {
            "/issues": [{"check": "TITLE_MISSING", "severity": "warning"}],
            "/pages": [],
        },
        binding,
    )

    from seohead.reports import build_report

    result = build_report(scan, fmt="md", path=str(tmp_path / "stale.md"))
    assert result["ok"] is False
    assert "warning issue count disagrees" in result["error"]
