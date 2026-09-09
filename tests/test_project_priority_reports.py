"""Human reports show saved priority decisions without changing audit or evidence."""

import csv
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from seohead.projects.coverage import coverage_status, initialize_coverage, record_execution
from seohead.projects.priorities import project_priorities
from seohead.projects.workspace import create_project
from seohead.reports import build_report
from tests.test_project_coverage_reports import AUDIT


@pytest.mark.parametrize("fmt", ["md", "csv", "xlsx", "docx"])
def test_saved_priorities_are_visible_and_read_only(tmp_path, fmt):
    root = tmp_path / "project"
    create_project(
        root,
        "https://example.test/",
        facts=[
            {"name": "framework", "value": "React", "provenance": "operator", "observed_at": None}
        ],
    )
    initialize_coverage(root)
    record_execution(
        root,
        "skill:workflow/js-render-check",
        {
            "status": "succeeded",
            "reason": "Reviewed synthetic rendering",
            "reviewer": "Specialist",
            "signoff": True,
        },
        expected_revision=1,
    )
    project_priorities(str(root), apply=True, expected_revision=2)
    status = coverage_status(root)
    row = next(r for r in status["items"] if r["id"] == "skill:workflow/js-render-check")
    assert row["priority"] == "P0" and row["complete"] and not row["stale"]
    before = (root / "coverage.json").read_bytes()
    target = tmp_path / f"report.{fmt}"
    result = build_report(AUDIT, fmt=fmt, path=str(target), project=str(root))
    assert result["ok"], result
    if fmt == "md":
        text = target.read_text()
        assert f"P0 (policy): {row['priority_reason']}" in text
    elif fmt == "csv":
        with target.with_suffix(".coverage.csv").open(encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream, delimiter=";"))
        actual = next(r for r in rows if r["Item ID"] == row["id"])
        assert (actual["Priority"], actual["Priority origin"], actual["Priority reason"]) == (
            "P0",
            "policy",
            row["priority_reason"],
        )
    elif fmt == "xlsx":
        rows = list(load_workbook(target)["Project Coverage"].values)
        actual = next(r for r in rows if r[0] == row["id"])
        assert actual[4:7] == ("P0", "policy", row["priority_reason"])
    else:
        text = "\n".join(p.text for p in Document(target).paragraphs)
        assert f"P0 (policy): {row['priority_reason']}" in text
    assert (root / "coverage.json").read_bytes() == before
    output = tmp_path / "original.json"
    assert build_report(AUDIT, fmt="json", path=str(output), project=str(root))["ok"]
    assert json.loads(output.read_text()) == AUDIT
