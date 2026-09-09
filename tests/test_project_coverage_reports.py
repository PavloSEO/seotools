"""Human reports preserve one project checklist snapshot without running it."""

from __future__ import annotations

import csv
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from seohead.audit.site import SCHEMA
from seohead.projects.coverage import (
    coverage_status,
    initialize_coverage,
    record_execution,
    update_item,
)
from seohead.projects.workspace import create_project
from seohead.reports import build_report

AUDIT = {
    "schema": SCHEMA,
    "domain": "example.test",
    "url": "https://example.test/",
    "generated_at": "2026-09-09T00:00:00Z",
    "findings": [],
    "pages": [],
    "summary": {"pages_checked": 1, "findings_total": 0, "findings_by_severity": {}},
}


@pytest.fixture
def project(tmp_path):
    root = tmp_path / "project"
    create_project(root, "https://example.test/")
    initialize_coverage(root)
    revision = coverage_status(root)["revision"]
    update_item(
        root,
        {"id": "custom:formula-review", "title": "=formula title"},
        expected_revision=revision,
    )
    revision = coverage_status(root)["revision"]
    record_execution(
        root,
        "custom:formula-review",
        {
            "status": "succeeded",
            "reason": "=formula reason",
            "reviewer": "Specialist",
            "signoff": True,
        },
        expected_revision=revision,
    )
    return root


@pytest.mark.parametrize("fmt", ["md", "csv", "xlsx", "docx"])
def test_human_reports_include_the_current_project_coverage_snapshot(project, tmp_path, fmt):
    target = tmp_path / f"audit.{fmt}"
    result = build_report(AUDIT, fmt=fmt, path=str(target), project=str(project))
    assert result["ok"], result
    expected = coverage_status(project)
    if fmt == "md":
        text = target.read_text(encoding="utf-8")
        assert "## Project checklist coverage" in text
        assert f"| {expected['counts']['total']} |" in text
        assert "=formula reason" in text
        assert "specialist review or signoff" in text
    elif fmt == "csv":
        coverage = target.with_suffix(".coverage.csv")
        assert coverage.exists() and str(coverage) in result["outputs"]
        rows = list(csv.DictReader(coverage.open(encoding="utf-8-sig"), delimiter=";"))
        formula = next(row for row in rows if row["Item ID"] == "custom:formula-review")
        assert formula["Counts"] == json.dumps(
            expected["counts"], ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        assert (
            formula["Measurement"]
            == '{"reason":"specialist review or signoff; no automatic measurement is implied","state":"not_measured"}'
        )
        assert formula["Reason"] == "'=formula reason"
    elif fmt == "xlsx":
        worksheet = load_workbook(target, data_only=False)["Project Coverage"]
        rows = list(worksheet.iter_rows(values_only=True))
        assert any(row[:2] == ("Checklist state", "initialized") for row in rows)
        formula = next(row for row in rows if row[0] == "custom:formula-review")
        assert formula[1] == "'=formula title"
        assert formula[-1] == "'=formula reason"
    else:
        text = "\n".join(paragraph.text for paragraph in Document(str(target)).paragraphs)
        assert "Project Checklist Coverage" in text
        assert "=formula reason" in text
        assert "specialist review or signoff" in text


def test_no_project_keeps_existing_human_outputs_and_json_passthrough_unchanged(tmp_path, project):
    markdown = tmp_path / "plain.md"
    plain = build_report(AUDIT, fmt="md", path=str(markdown))
    assert plain["ok"] and "Project checklist coverage" not in markdown.read_text(encoding="utf-8")
    csv_target = tmp_path / "plain.csv"
    csv_result = build_report(AUDIT, fmt="csv", path=str(csv_target))
    assert csv_result["ok"]
    assert not csv_target.with_suffix(".coverage.csv").exists()
    target = tmp_path / "audit.json"
    result = build_report(AUDIT, fmt="json", path=str(target), project=str(project))
    assert result["ok"]
    assert json.loads(target.read_text(encoding="utf-8")) == AUDIT


def test_project_report_refuses_a_different_audit_site(project, tmp_path):
    other = {**AUDIT, "domain": "other.test", "url": "https://other.test/"}
    result = build_report(other, fmt="md", path=str(tmp_path / "other.md"), project=str(project))
    assert result["ok"] is False
    assert "does not match audit source identity" in result["error"]


def test_project_report_binds_sf_audit_to_its_recorded_crawl_identity(project, tmp_path):
    audit = {
        "schema_version": "2.0",
        "run": {"source": "https://example.test/"},
        "summary": {"totals": {}, "by_severity": {}, "by_check": {}},
        "issues": [],
        "pages": [],
        "groups": [],
    }
    result = build_report(audit, fmt="md", path=str(tmp_path / "sf.md"), project=str(project))
    assert result["ok"], result


@pytest.mark.parametrize("fmt", ["md", "csv", "xlsx", "docx"])
def test_dependency_blocker_is_visible_when_a_completed_parent_becomes_stale(
    project, tmp_path, fmt
):
    revision = coverage_status(project)["revision"]
    update_item(project, {"id": "custom:parent", "title": "Parent"}, expected_revision=revision)
    revision = coverage_status(project)["revision"]
    update_item(
        project,
        {
            "id": "custom:child",
            "title": "Child\nreview",
            "dependencies": ["custom:parent"],
        },
        expected_revision=revision,
    )
    for item_id in ("custom:parent", "custom:child"):
        record_execution(
            project,
            item_id,
            {
                "status": "succeeded",
                "reason": "Reviewed",
                "reviewer": "Specialist",
                "signoff": True,
            },
            expected_revision=coverage_status(project)["revision"],
        )
    update_item(
        project,
        {"id": "custom:parent", "title": "Parent revised"},
        expected_revision=coverage_status(project)["revision"],
    )
    child = next(item for item in coverage_status(project)["items"] if item["id"] == "custom:child")
    assert child["state"] == "run" and child["complete"] is False
    assert child["blocked_by"] == ["custom:parent"]

    target = tmp_path / f"blocked.{fmt}"
    result = build_report(AUDIT, fmt=fmt, path=str(target), project=str(project))
    assert result["ok"], result
    if fmt == "md":
        text = target.read_text(encoding="utf-8")
        child_row = next(line for line in text.splitlines() if "Child review" in line)
        assert "Child\nreview" not in text
        assert "False" in child_row and '["custom:parent"]' in child_row
    elif fmt == "csv":
        rows = csv.DictReader(
            target.with_suffix(".coverage.csv").open(encoding="utf-8-sig"), delimiter=";"
        )
        child_row = next(row for row in rows if row["Item ID"] == "custom:child")
        assert child_row["Complete"] == "False"
        assert child_row["Blocked by"] == '["custom:parent"]'
    elif fmt == "xlsx":
        rows = list(load_workbook(target, data_only=False)["Project Coverage"].values)
        values = next(row for row in rows if row[0] == "custom:child")
        child_row = dict(zip(rows[6], values, strict=True))
        assert child_row["Complete"] is False and child_row["Blocked by"] == '["custom:parent"]'
    else:
        text = "\n".join(paragraph.text for paragraph in Document(str(target)).paragraphs)
        assert "Complete: False" in text
        assert 'Blocked by: ["custom:parent"]' in text


def test_project_report_never_replaces_controls_or_referenced_artifact_sidecars(project):
    control = build_report(
        AUDIT, fmt="md", path=str(project / "project.json"), project=str(project)
    )
    assert control["ok"] is False
    control = build_report(
        AUDIT, fmt="md", path=str(project / "coverage.json"), project=str(project)
    )
    assert control["ok"] is False
    evidence = project / "reports" / "protected.scope.csv"
    evidence.write_text("reviewed", encoding="utf-8")
    revision = coverage_status(project)["revision"]
    update_item(project, {"id": "custom:protected"}, expected_revision=revision)
    revision = coverage_status(project)["revision"]
    record_execution(
        project,
        "custom:protected",
        {
            "status": "succeeded",
            "reason": "Reviewed",
            "artifact": "reports/protected.scope.csv",
            "reviewer": "Specialist",
            "review": "approved",
        },
        expected_revision=revision,
    )
    collision = build_report(
        AUDIT, fmt="csv", path=str(project / "reports" / "protected.csv"), project=str(project)
    )
    assert collision["ok"] is False
    assert "referenced evidence" in collision["error"]
