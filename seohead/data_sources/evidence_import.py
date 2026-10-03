"""Offline import and normalization of analytics/search evidence for scan joins.

Provider rows (Google Search Console, GA4, Metrika, webmaster tools) and supplied
CSV/XLSX/JSON exports all raise the same questions before they may join a crawl:
which URL key does a row belong to, which metric is measured versus merely
absent, which reporting window, timezone, attribution, engine and identity
produced the numbers, and where each of those facts was established. This
module answers them once, in a versioned contract, so the normalized join in
``seohead.data_sources.evidence_join`` — and later the BI projection in #834
and the cohort quadrants in #835 — consume the same normalized observation
instead of re-deriving it per consumer.

Two artifacts:

- ``seohead.evidence-mapping.v1`` — a bounded, data-only manifest declaring
  how raw rows map onto dimensions, metrics, a URL field or dimension, an
  inclusive reporting period, timezone, attribution and collection coverage.
  Every resolved field records where it came from: ``envelope`` (a saved
  provider artifact), ``declared`` (the manifest or a documented provider
  contract), or ``unknown``. An unknown timezone, attribution, engine or
  reporting identity stays unknown — it is never fabricated.
- ``seohead.normalized-evidence.v1`` — one normalized row per input row at the
  declared grain. A supplied numeric zero stays a measured zero; an empty
  cell, a missing key, JSON ``null``, a non-numeric or formula cell, a
  suppressed placeholder and a never-collected row are each a distinct
  unavailable state, never zero.

Everything here is offline: importing a saved provider envelope or a supplied
export never calls a provider, resolves DNS, or fetches a page.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

from seohead.data_sources.providers import EVIDENCE_FORMAT
from seohead.tools.external_join import normalize_join_key

MAPPING_FORMAT = "seohead.evidence-mapping.v1"
NORMALIZED_FORMAT = "seohead.normalized-evidence.v1"

MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 256

COLLECTION_STATES = ("complete", "partial", "failed", "skipped", "not_configured")
# Collection coverage that still yields measured rows but must remain visible.
DEGRADED_COVERAGE = ("partial", "sampled", "thresholded", "truncated")
UNAVAILABLE_REASONS = (
    "absent",
    "empty",
    "null",
    "invalid",
    "non_finite",
    "formula",
    "suppressed",
)
ROW_SHAPES = ("flat", "gsc", "ga4")
FILE_KINDS = ("csv", "xlsx", "json", "provider_evidence")
URL_KINDS = ("absolute", "relative_path")
DUPLICATE_POLICIES = ("mark", "reject")

# Facts a provider contract guarantees for every response of an operation —
# Search Analytics reports Pacific-time calendar days for Google web search.
# These are declarations, not observations, so field provenance records them as
# ``declared`` with ``declared_by: "provider_contract"``. Providers absent here
# (GA4, Metrika) establish no timezone, engine or attribution by default.
_PROVIDER_CONTRACT: dict[str, dict[str, Any]] = {
    "gsc": {
        "search_engine": "google",
        "search_type": "web",
        "timezone": "America/Los_Angeles",
    },
    "bing_webmaster": {"search_engine": "bing", "search_type": "web"},
    "yandex_webmaster": {"search_engine": "yandex", "search_type": "web"},
}

# Metric cells carrying these placeholders are suppressed values — the source
# hid the number — which is a different fact from a parse failure.
_SUPPRESSED_STRINGS = frozenset({"(other)", "(not set)"})
# Dimension placeholders a provider emits instead of a real URL.
_NON_URL_STRINGS = frozenset({"(not set)", "(other)"})

_SOURCE_KEYS = {
    "kind",
    "provider",
    "operation",
    "reporting_identity",
    "privacy",
    "search_engine",
    "search_type",
    "attribution",
    "timezone",
    "site_origin",
}
_COLLECTION_KEYS = {
    "collected_at",
    "state",
    "sampled",
    "thresholded",
    "truncated",
    "reason",
}
_URL_KEYS = {"field", "dimension", "kind"}
_METRIC_KEYS = {"name", "type", "unit"}
_MANIFEST_KEYS = {
    "format",
    "source",
    "period",
    "url",
    "metrics",
    "dimensions",
    "collection",
    "duplicate_policy",
    "row_shape",
}
# Values that identify the reporting property or the client site; kept in the
# restricted artifact, replaced by a redacted marker in public responses.
_PRIVATE_FIELDS = frozenset({"reporting_identity", "site_origin"})
_ABSENT = object()


class EvidenceImportError(ValueError):
    """The supplied file or manifest cannot be normalized as declared."""


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _is_iso_date(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def _json_safe_cell(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


# --- file loading -------------------------------------------------------------


def _detect_kind(path: Path, raw: bytes) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix == ".xlsx":
        return "xlsx"
    if suffix == ".json":
        return "json"
    if suffix in {".jsonl", ".ndjson"}:
        raise EvidenceImportError("JSONL row streams are not a supported evidence shape")
    stripped = raw.lstrip()[:1]
    if stripped == b"{" or stripped == b"[":
        return "json"
    raise EvidenceImportError(
        f"cannot determine evidence kind for {path.name!r}; use .csv, .xlsx or .json"
    )


def _read_csv(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            columns = list(reader.fieldnames or [])
            if len(columns) > MAX_COLUMNS:
                raise EvidenceImportError(f"CSV exceeds the {MAX_COLUMNS}-column bound")
            rows: list[dict[str, Any]] = []
            for row in reader:
                rows.append({key: value for key, value in row.items() if key is not None})
                if len(rows) > MAX_ROWS:
                    raise EvidenceImportError(f"CSV exceeds the {MAX_ROWS}-row bound")
    except UnicodeDecodeError as exc:
        raise EvidenceImportError("CSV is not decodable as UTF-8") from exc
    return rows, columns


def _read_xlsx(path: Path, *, sheet: str | None) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl is a core dependency
        raise EvidenceImportError("XLSX input requires the openpyxl dependency") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise EvidenceImportError(f"cannot read XLSX workbook: {exc}") from exc
    try:
        names = workbook.sheetnames
        if sheet is not None:
            if sheet not in names:
                raise EvidenceImportError(
                    f"sheet {sheet!r} not found in workbook; available sheets: {names}"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise EvidenceImportError("XLSX sheet has no header row")
        columns = [
            str(cell) if cell is not None else f"column_{index + 1}"
            for index, cell in enumerate(header)
        ]
        if len(columns) > MAX_COLUMNS:
            raise EvidenceImportError(f"XLSX exceeds the {MAX_COLUMNS}-column bound")
        if len(set(columns)) != len(columns):
            raise EvidenceImportError("XLSX header contains duplicate column names")
        rows = []
        for raw_row in iterator:
            row = {
                columns[index]: _json_safe_cell(raw_row[index]) if index < len(raw_row) else None
                for index in range(len(columns))
            }
            if all(value is None for value in row.values()):
                continue  # a blank sheet row is not an observation
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise EvidenceImportError(f"XLSX exceeds the {MAX_ROWS}-row bound")
        return rows, columns
    finally:
        workbook.close()


def _read_json(raw: bytes) -> tuple[str, list[dict[str, Any]], dict | None, dict | None]:
    try:
        body = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise EvidenceImportError(f"JSON evidence is not parseable: {exc}") from exc
    if (
        isinstance(body, dict)
        and set(body) == {"evidence", "result"}
        and isinstance(body["evidence"], dict)
        and body["evidence"].get("format") == EVIDENCE_FORMAT
        and isinstance(body["result"], dict)
    ):
        result = body["result"]
        rows = result.get("rows")
        if (
            not isinstance(rows, list)
            or len(rows) > MAX_ROWS
            or any(not isinstance(row, dict) for row in rows)
        ):
            raise EvidenceImportError(
                "saved provider operation does not contain bounded joinable rows"
            )
        return "provider_evidence", rows, body["evidence"], result
    if isinstance(body, dict) and isinstance(body.get("rows"), list):
        rows = body["rows"]
    elif isinstance(body, list):
        rows = body
    else:
        raise EvidenceImportError(
            "JSON evidence must be a row array, an object with a rows array, "
            "or a saved seohead.provider-evidence.v1 envelope"
        )
    if len(rows) > MAX_ROWS or any(not isinstance(row, dict) for row in rows):
        raise EvidenceImportError("JSON evidence requires a bounded list of row objects")
    return "json", rows, None, None


def load_import_file(path: str | Path, *, sheet: str | None = None) -> dict[str, Any]:
    """Read one bounded supplied CSV/XLSX/JSON or saved provider envelope, offline.

    Returns ``{"kind", "rows", "columns", "envelope", "result", "file"}``.
    ``kind == "provider_evidence"`` means the file is a saved
    ``seohead.provider-evidence.v1`` envelope; its ``result`` then carries the
    raw provider rows and metadata. Formula cells in an XLSX are kept as their
    formula text — never evaluated and never guessed from a cached value.
    """
    source = Path(path)
    if source.is_symlink() or not source.is_file():
        raise EvidenceImportError("evidence file must be a regular file")
    size = source.stat().st_size
    if size > MAX_FILE_BYTES:
        raise EvidenceImportError(f"evidence file exceeds the {MAX_FILE_BYTES}-byte bound")
    raw = source.read_bytes()
    kind = _detect_kind(source, raw)
    file_meta = {"name": source.name, "sha256": _sha256(raw), "bytes": size}
    envelope: dict | None = None
    result: dict | None = None
    if kind == "csv":
        rows, columns = _read_csv(source)
    elif kind == "xlsx":
        rows, columns = _read_xlsx(source, sheet=sheet)
    else:
        kind, rows, envelope, result = _read_json(raw)
        columns = sorted({key for row in rows for key in row})[: MAX_COLUMNS + 1]
        if len(columns) > MAX_COLUMNS:
            raise EvidenceImportError(f"JSON evidence exceeds the {MAX_COLUMNS}-column bound")
    if kind == "provider_evidence":
        from seohead.data_sources import credentials

        if not credentials.is_private_mode(source.stat().st_mode):
            raise EvidenceImportError("saved provider evidence must live in a private-mode file")
    file_meta["kind"] = kind
    return {
        "kind": kind,
        "rows": rows,
        "columns": columns,
        "envelope": envelope,
        "result": result,
        "file": file_meta,
    }


# --- manifest validation and resolution ---------------------------------------


def _validated_origin(value: Any) -> str | None:
    """A URL-binding origin must be an absolute http(s) scheme + authority."""
    if value is None:
        return None
    if not isinstance(value, str):
        raise EvidenceImportError("site_origin must be an absolute http(s) origin string")
    try:
        parts = urlsplit(value.strip())
    except ValueError as exc:
        raise EvidenceImportError(f"site_origin is not a parseable origin: {value!r}") from exc
    if parts.scheme not in {"http", "https"} or not parts.netloc or parts.path not in {"", "/"}:
        raise EvidenceImportError(
            f"site_origin must be an absolute http(s) origin without a path, got {value!r}"
        )
    return f"{parts.scheme}://{parts.netloc}".lower()


def validate_manifest(manifest: Any) -> dict[str, Any]:
    """Check one declared mapping manifest; an absent manifest is not an error.

    Unknown keys are refused rather than ignored — a misspelled manifest field
    must not silently become a missing declaration.
    """
    if manifest is None:
        return {}
    if not isinstance(manifest, dict) or manifest.get("format") != MAPPING_FORMAT:
        raise EvidenceImportError(f"mapping requires format {MAPPING_FORMAT}")
    unknown = set(manifest) - _MANIFEST_KEYS
    if unknown:
        raise EvidenceImportError(f"mapping has unsupported keys: {sorted(unknown)}")
    source = manifest.get("source")
    if source is not None:
        if not isinstance(source, dict):
            raise EvidenceImportError("mapping source must be an object")
        unknown = set(source) - _SOURCE_KEYS
        if unknown:
            raise EvidenceImportError(f"mapping source has unsupported keys: {sorted(unknown)}")
        for name in (
            "kind",
            "provider",
            "operation",
            "reporting_identity",
            "privacy",
            "search_engine",
            "search_type",
            "attribution",
        ):
            if source.get(name) is not None and not isinstance(source[name], str):
                raise EvidenceImportError(f"mapping source {name} must be a string")
        if source.get("kind") is not None and source["kind"] not in FILE_KINDS:
            raise EvidenceImportError(f"mapping source kind must be one of {FILE_KINDS}")
        if source.get("privacy") is not None and source["privacy"] not in {
            "restricted",
            "supplied",
        }:
            raise EvidenceImportError("mapping source privacy must be restricted or supplied")
        timezone = source.get("timezone")
        if timezone is not None:
            if not isinstance(timezone, str):
                raise EvidenceImportError("mapping source timezone must be an IANA name string")
            try:
                ZoneInfo(timezone)
            except (KeyError, ValueError) as exc:
                raise EvidenceImportError(
                    f"mapping source timezone is not a known IANA name: {timezone!r}"
                ) from exc
        if source.get("site_origin") is not None:
            _validated_origin(source["site_origin"])
    period = manifest.get("period")
    if period is not None:
        if not isinstance(period, dict) or set(period) - {"start_date", "end_date"}:
            raise EvidenceImportError("mapping period must hold only start_date and end_date")
        start, end = period.get("start_date"), period.get("end_date")
        if (start is None) != (end is None):
            raise EvidenceImportError("mapping period requires both start_date and end_date")
        if start is not None and (not _is_iso_date(start) or not _is_iso_date(end) or start > end):
            raise EvidenceImportError(
                "mapping period must be an inclusive ordered pair of YYYY-MM-DD dates"
            )
    url_spec = manifest.get("url")
    if url_spec is not None:
        if not isinstance(url_spec, dict) or set(url_spec) - _URL_KEYS:
            raise EvidenceImportError("mapping url may name only field, dimension and kind")
        if url_spec.get("kind") is not None and url_spec["kind"] not in URL_KINDS:
            raise EvidenceImportError(f"mapping url kind must be one of {URL_KINDS}")
        for name in ("field", "dimension"):
            if url_spec.get(name) is not None and not isinstance(url_spec[name], str):
                raise EvidenceImportError(f"mapping url {name} must be a string")
    metrics = manifest.get("metrics")
    if metrics is not None:
        if not isinstance(metrics, list):
            raise EvidenceImportError("mapping metrics must be a list")
        for metric in metrics:
            if not isinstance(metric, dict) or set(metric) - _METRIC_KEYS:
                raise EvidenceImportError("each metric may name only name, type and unit")
            if not isinstance(metric.get("name"), str) or not metric["name"]:
                raise EvidenceImportError("each metric requires a non-empty name")
            if metric.get("type", "number") != "number":
                raise EvidenceImportError(
                    f"metric {metric['name']!r} has unsupported type {metric.get('type')!r}; "
                    "only numeric metrics join evidence"
                )
            if metric.get("unit") is not None and not isinstance(metric["unit"], str):
                raise EvidenceImportError("metric unit must be a string")
    dimensions = manifest.get("dimensions")
    if dimensions is not None and (
        not isinstance(dimensions, list)
        or any(not isinstance(name, str) or not name for name in dimensions)
    ):
        raise EvidenceImportError("mapping dimensions must be a list of non-empty names")
    collection = manifest.get("collection")
    if collection is not None:
        if not isinstance(collection, dict) or set(collection) - _COLLECTION_KEYS:
            raise EvidenceImportError(
                f"mapping collection may name only {sorted(_COLLECTION_KEYS)}"
            )
        state = collection.get("state")
        if state is not None and state not in COLLECTION_STATES:
            raise EvidenceImportError(
                f"mapping collection state must be one of {COLLECTION_STATES}"
            )
        for flag in ("sampled", "thresholded", "truncated"):
            if collection.get(flag) is not None and not isinstance(collection[flag], bool):
                raise EvidenceImportError(f"mapping collection {flag} must be a boolean")
    duplicate_policy = manifest.get("duplicate_policy")
    if duplicate_policy is not None and duplicate_policy not in DUPLICATE_POLICIES:
        raise EvidenceImportError(f"duplicate_policy must be one of {DUPLICATE_POLICIES}")
    row_shape = manifest.get("row_shape")
    if row_shape is not None and row_shape not in ROW_SHAPES:
        raise EvidenceImportError(f"row_shape must be one of {ROW_SHAPES}")
    return manifest


def _period_from(value: Any) -> dict[str, str] | None:
    """Normalize an evidence period into {start_date, end_date} or ``None``."""
    if isinstance(value, dict):
        start, end = value.get("start_date"), value.get("end_date")
    elif isinstance(value, str) and ".." in value:
        start, end = value.split("..", 1)
    elif _is_iso_date(value):
        start = end = value
    else:
        return None
    if _is_iso_date(start) and _is_iso_date(end) and start <= end:
        return {"start_date": start, "end_date": end}
    return None


def _flag(value: Any) -> bool | None:
    """Map an evidence coverage flag to True/False, keeping unknown as ``None``."""
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value in {"unknown", "", "none"}:
        return None
    return bool(value) if value is not None else None


def resolve_mapping(
    manifest: Any,
    *,
    kind: str,
    envelope: dict | None = None,
    result: dict | None = None,
    columns: list[str] | None = None,
    rows: list[dict[str, Any]] | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Merge a declared manifest with envelope facts into the resolved mapping.

    Returns ``(resolved, provenance)`` where ``resolved`` is the effective
    mapping used for normalization and ``provenance`` records, per field, the
    value and whether it came from the ``envelope``, was ``declared`` (by the
    manifest or a documented provider contract) or remains ``unknown``.
    """
    manifest = validate_manifest(manifest)
    declared = manifest.get("source") or {}
    envelope = envelope or {}
    result = result or {}
    fields: dict[str, dict[str, Any]] = {}

    def pick(
        name: str,
        manifest_value: Any,
        envelope_value: Any,
        contract_key: str | None = None,
    ) -> Any:
        provider = fields.get("provider", {}).get("value") or envelope.get("provider")
        contract = _PROVIDER_CONTRACT.get(provider or "", {})
        if manifest_value is not None:
            fields[name] = {
                "value": manifest_value,
                "origin": "declared",
                "declared_by": "manifest",
            }
            return manifest_value
        if envelope_value is not None:
            fields[name] = {"value": envelope_value, "origin": "envelope"}
            return envelope_value
        if contract_key and contract.get(contract_key) is not None:
            fields[name] = {
                "value": contract[contract_key],
                "origin": "declared",
                "declared_by": "provider_contract",
            }
            return contract[contract_key]
        fields[name] = {"value": None, "origin": "unknown"}
        return None

    provider = pick("provider", declared.get("provider"), envelope.get("provider"))
    operation = pick("operation", declared.get("operation"), envelope.get("operation"))
    reporting_identity = pick(
        "reporting_identity",
        declared.get("reporting_identity"),
        envelope.get("reporting_identity") or result.get("reporting_identity"),
    )
    site_origin = pick(
        "site_origin",
        _validated_origin(declared.get("site_origin")),
        envelope.get("site_origin"),
    )
    search_engine = pick(
        "search_engine",
        declared.get("search_engine"),
        envelope.get("search_engine"),
        "search_engine",
    )
    search_type = pick(
        "search_type",
        declared.get("search_type"),
        envelope.get("search_type"),
        "search_type",
    )
    attribution = pick(
        "attribution",
        declared.get("attribution"),
        envelope.get("attribution"),
    )
    timezone = pick(
        "timezone",
        declared.get("timezone"),
        envelope.get("timezone"),
        "timezone",
    )
    collected_at = pick(
        "collected_at",
        (manifest.get("collection") or {}).get("collected_at"),
        envelope.get("retrieved_at"),
    )
    collection_state = pick(
        "collection_state",
        (manifest.get("collection") or {}).get("state"),
        envelope.get("status"),
    )
    manifest_collection = manifest.get("collection") or {}
    sampled = pick(
        "sampled",
        manifest_collection.get("sampled"),
        _flag(envelope.get("sampling")),
    )
    thresholded = pick(
        "thresholded",
        manifest_collection.get("thresholded"),
        _flag(envelope.get("privacy_thresholds")),
    )
    truncated = pick(
        "truncated",
        manifest_collection.get("truncated"),
        _flag((envelope.get("pagination") or {}).get("truncated") or result.get("truncated")),
    )
    period = _period_from(manifest.get("period")) or _period_from(envelope.get("period"))
    if period is None:
        period = _period_from(result.get("period"))
    fields["period"] = {
        "value": period,
        "origin": (
            "declared"
            if _period_from(manifest.get("period"))
            else "envelope"
            if period is not None
            else "unknown"
        ),
    }
    if _period_from(manifest.get("period")):
        fields["period"]["declared_by"] = "manifest"

    privacy = declared.get("privacy") or (
        "restricted" if kind == "provider_evidence" or provider else "supplied"
    )
    fields["privacy"] = {
        "value": privacy,
        "origin": "declared",
        "declared_by": "manifest" if declared.get("privacy") else "import_default",
    }

    columns = columns or []
    row_shape = manifest.get("row_shape")
    if row_shape is None:
        sample = [row for row in (rows or result.get("rows") or [])[:50] if isinstance(row, dict)]
        if any("metricValues" in row for row in sample):
            row_shape = "ga4"
        elif any("keys" in row for row in sample):
            row_shape = "gsc"
        elif provider == "ga4":
            row_shape = "ga4"
        elif provider == "gsc":
            row_shape = "gsc"
        else:
            row_shape = "flat"

    url_spec = dict(manifest.get("url") or {})
    url_kind = url_spec.get("kind") or "absolute"
    url_field = url_spec.get("field")
    url_dimension = url_spec.get("dimension")
    dimensions = list(manifest.get("dimensions") or [])
    if not url_field and not url_dimension:
        envelope_dims = result.get("dimensions") or envelope.get("dimensions") or []
        if row_shape == "ga4" or "landingPagePlusQueryString" in envelope_dims:
            url_dimension = "landingPagePlusQueryString"
            url_kind = "relative_path"
        elif "page" in envelope_dims:
            url_dimension = "page"
        elif row_shape == "flat" and "url" in columns:
            url_field = "url"
        elif row_shape == "flat":
            raise EvidenceImportError("cannot determine the URL field; declare mapping url.field")
        else:
            raise EvidenceImportError(
                "cannot determine the URL dimension; declare mapping url.dimension"
            )
    if not dimensions:
        dimensions = list(result.get("dimensions") or envelope.get("dimensions") or [])
    envelope_dims = result.get("dimensions") or envelope.get("dimensions") or []
    fields["url_binding"] = {
        "value": {"field": url_field, "dimension": url_dimension, "kind": url_kind},
        "origin": (
            "declared"
            if manifest.get("url") or url_field
            else "envelope"
            if url_dimension in envelope_dims
            else "unknown"
        ),
        "declared_by": "manifest" if manifest.get("url") else "import_default",
    }

    metrics = manifest.get("metrics")
    if metrics is None:
        metrics = []
        if row_shape == "ga4":
            metrics = [
                {"name": name, "type": "number", "unit": "count"}
                for name in (result.get("metrics") or [])
            ]
        elif row_shape == "gsc":
            known = [m for m in ("clicks", "impressions", "ctr", "position")]
            rows = result.get("rows") or []
            present = {key for row in rows[:50] for key in row if key != "keys"}
            metrics = [
                {
                    "name": name,
                    "type": "number",
                    "unit": "ratio"
                    if name == "ctr"
                    else "position"
                    if name == "position"
                    else "count",
                }
                for name in known
                if name in present
            ]
    fields["metrics"] = {
        "value": [metric["name"] for metric in metrics],
        "origin": "declared"
        if manifest.get("metrics") is not None
        else "envelope"
        if metrics
        else "unknown",
        "declared_by": "manifest" if manifest.get("metrics") is not None else "import_default",
    }
    fields["dimensions"] = {
        "value": dimensions,
        "origin": "declared"
        if manifest.get("dimensions") is not None
        else "envelope"
        if dimensions
        else "unknown",
        "declared_by": "manifest" if manifest.get("dimensions") is not None else "import_default",
    }

    collection = {
        "collected_at": collected_at,
        "state": collection_state or "complete",
        "sampled": sampled,
        "thresholded": thresholded,
        "truncated": truncated,
        "reason": manifest_collection.get("reason"),
    }
    resolved = {
        "format": MAPPING_FORMAT,
        "source": {
            "kind": kind,
            "provider": provider,
            "operation": operation,
            "reporting_identity": reporting_identity,
            "privacy": privacy,
            "search_engine": search_engine,
            "search_type": search_type,
            "attribution": attribution,
            "timezone": timezone,
            "site_origin": site_origin,
        },
        "period": period,
        "url": {"field": url_field, "dimension": url_dimension, "kind": url_kind},
        "metrics": [
            {"name": m["name"], "type": "number", "unit": m.get("unit") or "count"} for m in metrics
        ],
        "dimensions": dimensions,
        "collection": collection,
        "duplicate_policy": manifest.get("duplicate_policy") or "mark",
        "row_shape": row_shape,
    }
    return resolved, {"fields": fields}


# --- row normalization ---------------------------------------------------------


def _lookup(row: dict[str, Any], name: str) -> Any:
    """Find a flat column by exact name, then by a provider's tail segment."""
    if name in row:
        return row[name]
    tail = name.split(":")[-1]
    if tail != name and tail in row:
        return row[tail]
    return _ABSENT


def _typed_metric(raw: Any) -> dict[str, Any]:
    """Type one metric cell; ``0`` is measured, absence is never zero."""
    entry: dict[str, Any] = {
        "value": None,
        "state": "unavailable",
        "reason": None,
        "raw": raw if raw is not _ABSENT else None,
    }
    if raw is _ABSENT:
        entry["reason"] = "absent"
        return entry
    if raw is None:
        entry["reason"] = "null"
        return entry
    if isinstance(raw, bool):
        entry["reason"] = "invalid"
        return entry
    if isinstance(raw, (int, float)):
        if math.isfinite(raw):
            return {"value": raw, "state": "measured", "reason": None, "raw": raw}
        entry["reason"] = "non_finite"
        return entry
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            entry["reason"] = "empty"
            return entry
        if text.startswith("="):
            entry["reason"] = "formula"
            return entry
        if text.startswith(("<", "≤")) or text.casefold() in _SUPPRESSED_STRINGS:
            entry["reason"] = "suppressed"
            return entry
        try:
            value = float(text)
        except ValueError:
            entry["reason"] = "invalid"
            return entry
        if not math.isfinite(value):
            entry["reason"] = "non_finite"
            return entry
        return {"value": value, "state": "measured", "reason": None, "raw": raw}
    entry["reason"] = "invalid"
    return entry


def _dimension_rows(row: dict[str, Any], shape: str, dimensions: list[str]) -> dict[str, Any]:
    """Pull declared dimension values out of one raw row at its provider grain."""
    if shape == "gsc":
        keys = row.get("keys")
        keys = keys if isinstance(keys, list) else []
        return {
            name: keys[index] if index < len(keys) else None
            for index, name in enumerate(dimensions)
        }
    if shape == "ga4":
        values = row.get("dimensionValues")
        values = values if isinstance(values, list) else []
        out = {}
        for index, name in enumerate(dimensions):
            cell = values[index] if index < len(values) else None
            out[name] = cell.get("value") if isinstance(cell, dict) else None
        return out
    return {
        name: (None if _lookup(row, name) is _ABSENT else _lookup(row, name)) for name in dimensions
    }


def _metric_rows(row: dict[str, Any], shape: str, metrics: list[dict[str, Any]]) -> dict[str, Any]:
    """Pull declared metric values out of one raw row."""
    out = {}
    if shape == "ga4":
        values = row.get("metricValues")
        values = values if isinstance(values, list) else []
        for index, metric in enumerate(metrics):
            cell = values[index] if index < len(values) else None
            raw = cell.get("value") if isinstance(cell, dict) else _ABSENT
            out[metric["name"]] = _typed_metric(raw)
        return out
    for metric in metrics:
        out[metric["name"]] = _typed_metric(_lookup(row, metric["name"]))
    return out


def _row_url(
    row: dict[str, Any],
    dims: dict[str, Any],
    url_spec: dict[str, Any],
    site_origin: str | None,
    shape: str,
) -> dict[str, Any]:
    """Resolve one row's URL into a strict join key with an explicit reason on failure."""
    url: dict[str, Any] = {
        "raw": None,
        "resolved": None,
        "normalized": None,
        "state": "unkeyable",
        "reason": None,
    }
    raw = (
        dims.get(url_spec["dimension"])
        if url_spec.get("dimension")
        else _lookup(row, url_spec["field"])
    )
    if raw is _ABSENT:
        url["reason"] = "url_absent"
        return url
    url["raw"] = raw
    if raw is None:
        url["reason"] = "url_absent"
        return url
    if not isinstance(raw, str):
        url["reason"] = "url_invalid"
        return url
    text = raw.strip()
    if not text:
        url["reason"] = "url_empty"
        return url
    if text.casefold() in _NON_URL_STRINGS:
        url["reason"] = "non_url_value"
        return url
    candidate = text
    if url_spec.get("kind") == "relative_path" and not urlsplit(text).scheme:
        if site_origin is None:
            url["reason"] = "missing_site_origin"
            return url
        candidate = site_origin + (text if text.startswith("/") else f"/{text}")
    url["resolved"] = candidate
    normalized = normalize_join_key(candidate)
    if normalized is None:
        url["reason"] = "url_invalid"
        return url
    url["normalized"] = normalized
    url["state"] = "keyed"
    return url


def normalize_rows(
    rows: list[dict[str, Any]],
    resolved: dict[str, Any],
    provenance: dict[str, Any],
    file_meta: dict[str, Any],
) -> dict[str, Any]:
    """Normalize raw rows into one ``seohead.normalized-evidence.v1`` document.

    Every row keeps its declared grain: a page+query row is never collapsed
    into the page row, and duplicate natural keys are explicitly marked
    (or refused under ``duplicate_policy: "reject"``) rather than summed
    or silently dropped.
    """
    source = resolved["source"]
    url_spec = resolved["url"]
    dimensions = resolved["dimensions"]
    shape = resolved["row_shape"]
    site_origin = source.get("site_origin")
    url_dimension = url_spec.get("dimension")
    extra_dimensions = [
        name for name in dimensions if name not in {url_dimension, url_spec.get("field")}
    ]
    seen: dict[tuple, int] = {}
    ambiguous_indexes: set[int] = set()
    normalized_rows: list[dict[str, Any]] = []
    unkeyable_reasons: dict[str, int] = {}
    metric_summary = {
        metric["name"]: {"measured": 0, "unavailable": 0, "reasons": {}}
        for metric in resolved["metrics"]
    }
    keyed = 0
    for index, row in enumerate(rows):
        dims = _dimension_rows(row, shape, dimensions)
        url = _row_url(row, dims, url_spec, site_origin, shape)
        metrics = _metric_rows(row, shape, resolved["metrics"])
        for name, metric in metrics.items():
            bucket = metric_summary[name]
            if metric["state"] == "measured":
                bucket["measured"] += 1
            else:
                bucket["unavailable"] += 1
                bucket["reasons"][metric["reason"]] = bucket["reasons"].get(metric["reason"], 0) + 1
        natural = [url["raw"]] + [dims.get(name) for name in extra_dimensions]
        natural_key = tuple(json.dumps(value, default=str, ensure_ascii=False) for value in natural)
        if natural_key in seen:
            ambiguous_indexes.add(seen[natural_key])
            ambiguous_indexes.add(index)
        else:
            seen[natural_key] = index
        if url["state"] == "keyed":
            keyed += 1
        else:
            unkeyable_reasons[url["reason"]] = unkeyable_reasons.get(url["reason"], 0) + 1
        normalized_rows.append(
            {
                "row_index": index,
                "natural_key": natural,
                "natural_key_sha256": _sha256(
                    json.dumps(natural, default=str, ensure_ascii=False).encode()
                ),
                "dimensions": dims,
                "url": url,
                "metrics": {
                    name: {**metric, "unit": spec["unit"]}
                    for (name, metric), spec in zip(
                        metrics.items(), resolved["metrics"], strict=True
                    )
                },
                "ambiguous": False,
                "raw": row,
            }
        )
    if ambiguous_indexes and resolved["duplicate_policy"] == "reject":
        raise EvidenceImportError(
            f"{len(ambiguous_indexes)} rows share a duplicate natural key at the "
            "declared grain; duplicate_policy is reject"
        )
    for index in ambiguous_indexes:
        normalized_rows[index]["ambiguous"] = True
    document = {
        "format": NORMALIZED_FORMAT,
        "mapping": resolved,
        "provenance": {
            "fields": provenance["fields"],
            "file": file_meta,
            "artifact_reference": None,
        },
        "rows": normalized_rows,
        "summary": {
            "rows": len(normalized_rows),
            "keyed": keyed,
            "unkeyable": len(normalized_rows) - keyed,
            "unkeyable_reasons": unkeyable_reasons,
            "ambiguous": len(ambiguous_indexes),
            "metrics": metric_summary,
            "collection": resolved["collection"],
        },
    }
    return document


def normalize_file(
    path: str | Path,
    *,
    manifest: dict[str, Any] | None = None,
    sheet: str | None = None,
    site_origin: str | None = None,
) -> dict[str, Any]:
    """Load a supplied file or saved provider envelope and normalize it.

    ``site_origin`` is the explicit URL binding for providers whose rows carry
    relative keys (GA4 ``landingPagePlusQueryString``); without it those rows
    stay unkeyable with reason ``missing_site_origin`` instead of being
    guessed against a fabricated host.
    """
    loaded = load_import_file(path, sheet=sheet)
    manifest = validate_manifest(manifest) if manifest else {}
    if site_origin is not None:
        manifest = dict(manifest)
        manifest["format"] = MAPPING_FORMAT
        manifest["source"] = {
            **(manifest.get("source") or {}),
            "site_origin": site_origin,
        }
        manifest = validate_manifest(manifest)
    resolved, provenance = resolve_mapping(
        manifest or None,
        kind=loaded["kind"],
        envelope=loaded["envelope"],
        result=loaded["result"],
        columns=loaded["columns"],
        rows=loaded["rows"],
    )
    document = normalize_rows(loaded["rows"], resolved, provenance, loaded["file"])
    envelope = loaded["envelope"] or {}
    document["provenance"]["artifact_reference"] = envelope.get("artifact_reference")
    return document


def normalize_inline(
    rows: list[dict[str, Any]],
    *,
    manifest: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Normalize an already-parsed row list against a declared manifest."""
    if (
        not isinstance(rows, list)
        or len(rows) > MAX_ROWS
        or any(not isinstance(row, dict) for row in rows)
    ):
        raise EvidenceImportError("rows must be a bounded list of row objects")
    manifest = validate_manifest(manifest) if manifest else {}
    kind = (manifest.get("source") or {}).get("kind") or "json"
    columns = sorted({key for row in rows for key in row})
    resolved, provenance = resolve_mapping(manifest or None, kind=kind, columns=columns, rows=rows)
    return normalize_rows(
        rows,
        resolved,
        provenance,
        {"name": "inline", "sha256": None, "bytes": None, "kind": kind},
    )


def public_provenance(provenance: dict[str, Any]) -> dict[str, Any]:
    """Provenance with private reporting identifiers replaced by redaction markers."""
    fields = {}
    for name, entry in (provenance.get("fields") or {}).items():
        if name in _PRIVATE_FIELDS and entry.get("value") is not None:
            fields[name] = {**entry, "value": None, "redacted": True}
        else:
            fields[name] = entry
    file_meta = {**(provenance.get("file") or {})}
    file_meta.pop("name", None)
    return {
        "fields": fields,
        "file": file_meta,
        "artifact_reference": provenance.get("artifact_reference"),
    }
