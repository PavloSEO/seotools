"""Read explicit list-mode URLs from ordinary exchange files.

The reader deliberately scans every text cell/value rather than assuming a
particular column name or export shape.  A redirect map, an analytics export,
or a sitemap fragment can therefore be checked without first reshaping it.
It only extracts absolute HTTP(S) URLs and preserves encounter order; list
mode keeps its established first-occurrence de-duplication when it fetches
those inputs.
"""

from __future__ import annotations

import csv
import re
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from defusedxml import ElementTree

_URL = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
_TRAILING_PUNCTUATION = ".,;:!]}"
_SUPPORTED_SUFFIXES = {".txt", ".csv", ".xlsx", ".xml"}


class UrlListError(ValueError):
    """An explicit URL-list file cannot be read as a supported input."""


def _urls_in(value: Any) -> Iterator[str]:
    if not isinstance(value, str):
        return
    for raw in _URL.findall(value):
        candidate = raw.rstrip(_TRAILING_PUNCTUATION)
        parts = urlsplit(candidate)
        if parts.scheme.lower() in {"http", "https"} and parts.hostname:
            yield candidate


def _text_values(path: Path) -> Iterable[str]:
    return path.read_text(encoding="utf-8-sig").splitlines()


def _csv_values(path: Path) -> Iterator[str]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.reader(handle):
            yield from row


def _xlsx_values(path: Path) -> Iterator[Any]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise UrlListError(f"could not read spreadsheet {path}: {exc}") from exc
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                yield from row
    finally:
        workbook.close()


def _xml_values(path: Path) -> Iterator[str]:
    try:
        for event, element in ElementTree.iterparse(path, events=("start", "end")):
            if event == "start":
                yield from element.attrib.values()
            else:
                if element.text:
                    yield element.text
                if element.tail:
                    yield element.tail
                element.clear()
    except (OSError, ElementTree.ParseError) as exc:
        raise UrlListError(f"could not read XML {path}: {exc}") from exc


def read_url_list(path: str | Path) -> list[str]:
    """Extract absolute HTTP(S) URLs from TXT, CSV, XLSX, or XML in source order.

    The file extension is intentionally the format choice: guessing a CSV as
    plain text would change quoting semantics, and treating an unsupported
    extension as text can conceal an operator typo.  No URL normalisation or
    deduplication happens here because list-mode collection already owns that
    established request-identity policy.
    """
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in _SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(_SUPPORTED_SUFFIXES))
        raise UrlListError(f"unsupported URL-list file {source}; expected one of {supported}")
    if not source.is_file():
        raise UrlListError(f"URL-list file does not exist or is not a file: {source}")

    if suffix == ".txt":
        values = _text_values(source)
    elif suffix == ".csv":
        values = _csv_values(source)
    elif suffix == ".xlsx":
        values = _xlsx_values(source)
    else:
        values = _xml_values(source)

    urls = [url for value in values for url in _urls_in(value)]
    if not urls:
        raise UrlListError(f"URL-list file contains no absolute HTTP(S) URLs: {source}")
    return urls
